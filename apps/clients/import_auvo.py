"""
Importação assistida de clientes do Auvo — LocusHub, 08/09/2026 (relatório
aprovado pelo usuário na mesma data). Upload → revisão/curadoria → confirmação,
exatamente o mesmo fluxo de três telas de `apps.equipment.legacy_import`
(especificação, seção 13, fluxo D) — restrito a Administrador
(`apps.accounts.permissions.CAN_IMPORT_CLIENTS`).

Este módulo só lê e classifica — nada é gravado no banco aqui. A decisão de
"o que fazer com cada linha" (criar/pular) é sempre humana, na tela de
revisão, e revalidada de novo no momento da confirmação
(`apps.clients.views_import`) — mesma disciplina de "nunca confiar só na
prévia da sessão" já usada pelo importador de equipamentos.

Layout esperado (exportação real do Auvo, homologada com
`clientes_auvo_08_09_2026.xlsx`, 716 clientes): cabeçalho na linha 1, uma
linha explicativa fixa na linha 2 (sempre ignorada — é assim que o próprio
Auvo gera o arquivo, não uma heurística), dados a partir da linha 3. As
colunas podem vir em qualquer ordem — resolvidas pelo nome do cabeçalho, não
pela posição.
"""

import io
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from difflib import SequenceMatcher

from openpyxl import load_workbook

from apps.clients.validators import is_valid_cnpj, is_valid_cpf, normalize_document

logger = logging.getLogger(__name__)

# Todas as colunas que este importador efetivamente lê (mapeamento completo,
# relatório de 08/09/2026) — usadas também para validar que o arquivo
# enviado é mesmo uma exportação de clientes do Auvo neste layout, antes de
# processar qualquer linha.
REQUIRED_HEADERS = [
    "Código",
    "Código externo",
    "Nome",
    "CPF ou CNPJ",
    "Razão social",
    "Endereço",
    "Complemento do Endereço",
    "Telefone corporativo",
    "E-mail corporativo",
    "Falar com",
    "Segmento",
    "Observação",
    "Observação Interna",
    "Status",
    "Anotação",
    "Data de cadastro",
    "Contribuinte do ICMS",
    "Inscrição Estadual",
    "Inscrição Municipal",
    "E-mail de cobrança",
    "CEP de cobrança",
    "Logradouro endereço de cobrança",
    "Número endereço de cobrança",
    "Complemento endereço de cobrança",
    "Bairro de cobrança",
    "Cidade de cobrança",
    "Estado de cobrança",
]

# Linhas de dados começam sempre na 3ª linha da planilha — linha 1 é
# cabeçalho, linha 2 é a linha explicativa fixa que o próprio Auvo inclui em
# toda exportação (confirmado na planilha real de homologação).
FIRST_DATA_ROW = 3

# Limiar de semelhança (0-1, `difflib.SequenceMatcher`) para classificar dois
# nomes/razões sociais como "possivelmente o mesmo cliente" quando nenhum dos
# dois tem CPF/CNPJ para comparar. Deliberadamente alto: o objetivo aqui é
# nunca deixar passar batido um duplicado óbvio, não é sugerir "parecidos" —
# um falso positivo vira um cliente legítimo bloqueado de entrar como NOVO,
# então erramos para o lado conservador (ver decisão do usuário, 08/09/2026:
# "Se não houver sinal relevante de duplicidade, permitir como NOVO").
NAME_MATCH_THRESHOLD = 0.92

NOVO = "NOVO"
JA_EXISTENTE = "JA_EXISTENTE"
POSSIVEL_DUPLICADO = "POSSIVEL_DUPLICADO"
INVALIDO = "INVALIDO"


class ClientImportError(Exception):
    """Erro que impede a leitura da planilha (arquivo errado, coluna faltando, etc.)."""


def _text(value) -> str:
    return str(value).strip() if value not in (None, "") else ""


def _normalize_name(value: str) -> str:
    """Só para comparação de semelhança — não é o valor salvo em lugar nenhum."""
    decomposed = unicodedata.normalize("NFKD", value)
    without_accents = "".join(c for c in decomposed if not unicodedata.combining(c))
    collapsed = re.sub(r"\s+", " ", without_accents).strip().upper()
    return re.sub(r"[^\w\s]", "", collapsed)


def _normalize_cep(value: str) -> str:
    """
    O CEP de cobrança vem da planilha real em pelo menos três formatos
    diferentes ("87.083-304", "87075-800", "86990000" — confirmado na
    homologação com a planilha real de 08/09/2026), e o formato com ponto
    extra tem 10 caracteres — mais do que `Address.cep` comporta
    (`max_length=9`, mesmo campo usado pelo cadastro manual, que sempre
    veio no formato padrão "00000-000"). Normaliza para esse único formato
    canônico sempre que os 8 dígitos estão presentes; mantém o texto como
    veio (truncado defensivamente) só no caso raro de não ter exatamente 8
    dígitos, para nunca estourar o campo do banco.
    """
    digits = re.sub(r"\D", "", value)
    if len(digits) == 8:
        return f"{digits[:5]}-{digits[5:]}"
    return value.strip()[:9]


def _split_multi(value: str) -> list[str]:
    """
    Auvo separa telefones/e-mails múltiplos por ' ; ' — normaliza espaços,
    remove vazios e duplicatas (case-insensitive), preservando a ordem do
    primeiro valor encontrado (decisão do usuário, 08/09/2026: "normalizar
    espaços; ignorar vazios; evitar duplicidade do mesmo telefone/e-mail").
    """
    seen: set[str] = set()
    result: list[str] = []
    for raw in value.split(";"):
        item = raw.strip()
        if not item:
            continue
        key = item.lower()
        if key in seen:
            continue
        seen.add(key)
        result.append(item)
    return result


def _header_index_map(header_row) -> dict[str, int]:
    found = {}
    for idx, cell in enumerate(header_row):
        if cell.value:
            found[str(cell.value).strip()] = idx
    return found


def _cell(row, index_map: dict[str, int], header: str):
    idx = index_map.get(header)
    if idx is None:
        return None
    return row[idx].value


@dataclass
class NewClientPayload:
    """Formato serializável (JSON-friendly) dos dados prontos para `create_client()`."""

    client_type: str
    company_name: str = ""
    document: str = ""
    trade_name: str = ""
    registration_status: str = ""
    phone: str = ""
    email: str = ""
    contact_name: str = ""
    notes: str = ""
    auvo_code: str = ""
    external_code: str = ""
    state_registration: str = ""
    municipal_registration: str = ""
    icms_taxpayer: str = ""
    billing_email: str = ""
    fiscal_cep: str = ""
    fiscal_logradouro: str = ""
    fiscal_numero: str = ""
    fiscal_complemento: str = ""
    fiscal_bairro: str = ""
    fiscal_cidade: str = ""
    fiscal_uf: str = ""


@dataclass
class ParsedClientRow:
    row_number: int
    classification: str
    reason: str = ""

    # Campos de exibição na tela de revisão.
    nome: str = ""
    razao_social: str = ""
    documento_display: str = ""
    cidade_uf: str = ""
    auvo_code: str = ""

    # Só preenchido/relevante quando classification == NOVO.
    payload: NewClientPayload | None = None

    @property
    def has_billing_address(self) -> bool:
        return self.payload is not None and bool(
            self.payload.fiscal_cep
            or self.payload.fiscal_logradouro
            or self.payload.fiscal_bairro
            or self.payload.fiscal_cidade
        )

    def to_session_dict(self) -> dict:
        return {
            "row_number": self.row_number,
            "classification": self.classification,
            "reason": self.reason,
            "nome": self.nome,
            "razao_social": self.razao_social,
            "documento_display": self.documento_display,
            "cidade_uf": self.cidade_uf,
            "auvo_code": self.auvo_code,
            "payload": self.payload.__dict__ if self.payload else None,
        }

    @classmethod
    def from_session_dict(cls, data: dict) -> "ParsedClientRow":
        return cls(
            row_number=data["row_number"],
            classification=data["classification"],
            reason=data["reason"],
            nome=data["nome"],
            razao_social=data["razao_social"],
            documento_display=data["documento_display"],
            cidade_uf=data["cidade_uf"],
            auvo_code=data["auvo_code"],
            payload=NewClientPayload(**data["payload"]) if data["payload"] else None,
        )


@dataclass
class _DedupContext:
    """Estado de deduplicação — visto contra o banco E contra linhas já aceitas nesta mesma planilha."""

    existing_documents: set[str]
    existing_auvo_codes: set[str]
    # (nome_normalizado, rótulo_para_exibir, cliente_pk_ou_None) — cliente_pk é
    # None para candidatos aceitos dentro desta própria planilha (ainda sem pk).
    name_pool: list[tuple[str, str, int | None]] = field(default_factory=list)

    def best_name_match(self, candidate_names: list[str]) -> tuple[str, int | None] | None:
        best_ratio = 0.0
        best: tuple[str, int | None] | None = None
        candidates_norm = [_normalize_name(n) for n in candidate_names if n]
        for norm_name, label, pk in self.name_pool:
            for candidate_norm in candidates_norm:
                if not norm_name or not candidate_norm:
                    continue
                ratio = SequenceMatcher(None, norm_name, candidate_norm).ratio()
                if ratio > best_ratio:
                    best_ratio, best = ratio, (label, pk)
        if best_ratio >= NAME_MATCH_THRESHOLD:
            return best
        return None

    def register_accepted(self, names: list[str], label: str) -> None:
        for name in names:
            if name:
                self.name_pool.append((_normalize_name(name), label, None))


def _build_dedup_context() -> _DedupContext:
    from apps.clients.models import Client

    existing_documents = set(Client.objects.exclude(document="").values_list("document", flat=True))
    existing_auvo_codes = set(Client.objects.exclude(auvo_code="").values_list("auvo_code", flat=True))

    name_pool: list[tuple[str, str, int | None]] = []
    for pk, company_name, trade_name in Client.objects.values_list("pk", "company_name", "trade_name"):
        label = company_name or trade_name or f"cliente #{pk}"
        for name in (company_name, trade_name):
            if name:
                name_pool.append((_normalize_name(name), label, pk))

    return _DedupContext(
        existing_documents=existing_documents, existing_auvo_codes=existing_auvo_codes, name_pool=name_pool
    )


def _build_notes(
    *,
    endereco_original: str = "",
    telefones_adicionais: list[str] | None = None,
    emails_adicionais: list[str] | None = None,
    segmento: str = "",
    observacao: str = "",
    observacao_interna: str = "",
    anotacao: str = "",
    data_cadastro: str = "",
) -> str:
    """
    Monta `Client.notes` de forma organizada, com rótulos claros, incluindo
    só o que tiver valor (decisão do usuário, 08/09/2026: "não incluir campos
    vazios", "cuidado para não transformar o campo em um dump da planilha").
    """
    lines: list[str] = []
    if endereco_original:
        lines.append(f"Endereço original (Auvo): {endereco_original}")
    if telefones_adicionais:
        lines.append(f"Telefones adicionais (Auvo): {'; '.join(telefones_adicionais)}")
    if emails_adicionais:
        lines.append(f"E-mails adicionais (Auvo): {'; '.join(emails_adicionais)}")
    if segmento:
        lines.append(f"Segmento (Auvo): {segmento}")
    if observacao:
        lines.append(f"Observação (Auvo): {observacao}")
    if observacao_interna:
        lines.append(f"Observação interna (Auvo): {observacao_interna}")
    if anotacao:
        lines.append(f"Anotação (Auvo): {anotacao}")
    if data_cadastro:
        lines.append(f"Data de cadastro no Auvo: {data_cadastro}")

    if not lines:
        return ""
    return "[Dados importados do Auvo]\n" + "\n".join(lines)


def parse_client_workbook(uploaded_file) -> list[ParsedClientRow]:
    """
    Lê a planilha de exportação de clientes do Auvo e devolve uma linha
    classificada (NOVO/JÁ EXISTENTE/POSSÍVEL DUPLICADO/INVÁLIDO) por
    cliente, pronta para a tela de revisão. Nada é gravado no banco aqui.
    """
    try:
        wb = load_workbook(io.BytesIO(uploaded_file.read()), data_only=True)
    except Exception as exc:  # openpyxl levanta várias exceções diferentes para arquivo inválido
        # O texto de `exc` é interno do openpyxl/Python (em inglês, às vezes
        # com caminho de arquivo/traceback) — nunca deve chegar ao usuário
        # final; fica só no log técnico (auditoria de idioma, ago/2026).
        logger.warning("Falha ao ler planilha de importação de clientes: %s", exc, exc_info=True)
        raise ClientImportError(
            "Não foi possível ler o arquivo como planilha Excel. Verifique se o arquivo não está "
            "corrompido e se é realmente um .xlsx."
        ) from exc

    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(min_row=1)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ClientImportError("A planilha está vazia.") from None
    index_map = _header_index_map(header_row)

    missing_headers = [h for h in REQUIRED_HEADERS if h not in index_map]
    if missing_headers:
        raise ClientImportError(
            "A planilha não tem as colunas esperadas de uma exportação de clientes do Auvo: "
            + ", ".join(missing_headers)
            + ". Confira se o arquivo é realmente a exportação de clientes (não outra planilha do Auvo)."
        )

    dedup = _build_dedup_context()
    parsed_rows: list[ParsedClientRow] = []

    for row_number, row in enumerate(ws.iter_rows(min_row=FIRST_DATA_ROW), start=FIRST_DATA_ROW):
        codigo = _text(_cell(row, index_map, "Código"))
        codigo_externo = _text(_cell(row, index_map, "Código externo"))
        nome = _text(_cell(row, index_map, "Nome"))
        documento_raw = _text(_cell(row, index_map, "CPF ou CNPJ"))
        razao_social = _text(_cell(row, index_map, "Razão social"))
        endereco = _text(_cell(row, index_map, "Endereço"))
        complemento_endereco = _text(_cell(row, index_map, "Complemento do Endereço"))
        telefone = _text(_cell(row, index_map, "Telefone corporativo"))
        email = _text(_cell(row, index_map, "E-mail corporativo"))
        falar_com = _text(_cell(row, index_map, "Falar com"))
        segmento = _text(_cell(row, index_map, "Segmento"))
        observacao = _text(_cell(row, index_map, "Observação"))
        observacao_interna = _text(_cell(row, index_map, "Observação Interna"))
        status = _text(_cell(row, index_map, "Status"))
        anotacao = _text(_cell(row, index_map, "Anotação"))
        data_cadastro = _text(_cell(row, index_map, "Data de cadastro"))
        icms = _text(_cell(row, index_map, "Contribuinte do ICMS"))
        inscricao_estadual = _text(_cell(row, index_map, "Inscrição Estadual"))
        inscricao_municipal = _text(_cell(row, index_map, "Inscrição Municipal"))
        email_cobranca = _text(_cell(row, index_map, "E-mail de cobrança"))
        cep_cobranca = _text(_cell(row, index_map, "CEP de cobrança"))
        logradouro_cobranca = _text(_cell(row, index_map, "Logradouro endereço de cobrança"))
        numero_cobranca = _text(_cell(row, index_map, "Número endereço de cobrança"))
        complemento_cobranca = _text(_cell(row, index_map, "Complemento endereço de cobrança"))
        bairro_cobranca = _text(_cell(row, index_map, "Bairro de cobrança"))
        cidade_cobranca = _text(_cell(row, index_map, "Cidade de cobrança"))
        estado_cobranca = _text(_cell(row, index_map, "Estado de cobrança"))

        # Linha completamente vazia (sobra de formatação no fim da planilha) — ignorar de vez,
        # não é INVÁLIDO, é ausência de dado.
        if not any([codigo, nome, razao_social, documento_raw]):
            continue

        documento_digits = normalize_document(documento_raw)
        cidade_uf = f"{cidade_cobranca}/{estado_cobranca}" if (cidade_cobranca or estado_cobranca) else ""

        # --- Validação do documento -------------------------------------
        client_type = "PJ"
        invalid_reason = ""
        if documento_digits:
            if len(documento_digits) == 11:
                client_type = "PF"
                if not is_valid_cpf(documento_digits):
                    invalid_reason = "CPF inválido (dígito verificador não confere)."
            elif len(documento_digits) == 14:
                client_type = "PJ"
                if not is_valid_cnpj(documento_digits):
                    invalid_reason = "CNPJ inválido (dígito verificador não confere)."
            else:
                invalid_reason = f"CPF/CNPJ com {len(documento_digits)} dígito(s) — nem CPF (11) nem CNPJ (14)."

        if not invalid_reason and not nome and not razao_social:
            invalid_reason = "Nome e razão social ausentes."

        base_row_kwargs = dict(
            row_number=row_number,
            nome=nome,
            razao_social=razao_social,
            documento_display=documento_digits or "(sem documento)",
            cidade_uf=cidade_uf,
            auvo_code=codigo,
        )

        if invalid_reason:
            parsed_rows.append(ParsedClientRow(classification=INVALIDO, reason=invalid_reason, **base_row_kwargs))
            continue

        # --- Deduplicação --------------------------------------------------
        if codigo and codigo in dedup.existing_auvo_codes:
            parsed_rows.append(
                ParsedClientRow(
                    classification=JA_EXISTENTE,
                    reason=f"Código Auvo {codigo} já foi importado anteriormente.",
                    **base_row_kwargs,
                )
            )
            continue

        if documento_digits and documento_digits in dedup.existing_documents:
            parsed_rows.append(
                ParsedClientRow(
                    classification=JA_EXISTENTE,
                    reason=f"CPF/CNPJ {documento_digits} já cadastrado no LocusHub.",
                    **base_row_kwargs,
                )
            )
            continue

        if not documento_digits:
            match = dedup.best_name_match([nome, razao_social])
            if match is not None:
                matched_label, matched_pk = match
                reason = (
                    f"Cliente sem documento com nome/razão social semelhante a '{matched_label}'"
                    + (f" (cliente #{matched_pk})" if matched_pk else " (outra linha desta mesma planilha)")
                    + "."
                )
                parsed_rows.append(ParsedClientRow(classification=POSSIVEL_DUPLICADO, reason=reason, **base_row_kwargs))
                continue

        # --- NOVO: monta o payload pronto para create_client() -------------
        telefones = _split_multi(telefone)
        emails = _split_multi(email)
        phone, telefones_extra = (telefones[0], telefones[1:]) if telefones else ("", [])
        client_email, emails_extra = (emails[0], emails[1:]) if emails else ("", [])

        has_structured_billing_address = bool(cep_cobranca or logradouro_cobranca or bairro_cobranca or cidade_cobranca)
        endereco_original_para_notes = ""
        if not has_structured_billing_address and endereco:
            endereco_original_para_notes = endereco
            if complemento_endereco:
                endereco_original_para_notes = f"{endereco_original_para_notes} — {complemento_endereco}"

        notes = _build_notes(
            endereco_original=endereco_original_para_notes,
            telefones_adicionais=telefones_extra,
            emails_adicionais=emails_extra,
            segmento=segmento,
            observacao=observacao,
            observacao_interna=observacao_interna,
            anotacao=anotacao,
            data_cadastro=data_cadastro,
        )

        payload = NewClientPayload(
            client_type=client_type,
            company_name=razao_social,
            document=documento_digits,
            trade_name=nome,
            registration_status=status,
            phone=phone,
            email=client_email,
            contact_name=falar_com,
            notes=notes,
            auvo_code=codigo,
            external_code=codigo_externo,
            state_registration=inscricao_estadual,
            municipal_registration=inscricao_municipal,
            icms_taxpayer=icms,
            billing_email=email_cobranca,
            fiscal_cep=_normalize_cep(cep_cobranca) if cep_cobranca else "",
            fiscal_logradouro=logradouro_cobranca if has_structured_billing_address else "",
            fiscal_numero=numero_cobranca,
            fiscal_complemento=complemento_cobranca,
            fiscal_bairro=bairro_cobranca,
            fiscal_cidade=cidade_cobranca,
            fiscal_uf=estado_cobranca,
        )

        parsed_rows.append(ParsedClientRow(classification=NOVO, payload=payload, **base_row_kwargs))

        # Linhas sem documento aceitas como NOVO entram no "pool" de nomes
        # para pegar duplicidade DENTRO da própria planilha (duas linhas
        # parecidas sem CPF/CNPJ não podem as duas virar NOVO).
        if not documento_digits:
            dedup.register_accepted([nome, razao_social], razao_social or nome)

    return parsed_rows
